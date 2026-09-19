"""Tests for generated-format exporters and the ZIP writer."""

from __future__ import annotations

import json
import zipfile
from io import BytesIO

import pytest

from app.export.diagrams import (
    ArchitectureLayer,
    render_architecture_png,
    render_architecture_text,
)
from app.export.errors import ExportLimitError
from app.export.exporters.document_exporters import (
    CsvExporter,
    DocxExporter,
    HtmlExporter,
    JsonExporter,
    MarkdownExporter,
    PdfExporter,
    TextExporter,
)
from app.export.exporters.spreadsheet_exporters import XlsxExporter
from app.export.exporters.zip_writer import (
    ZipEntry,
    build_manifest,
    safe_relpath,
    write_zip_to_path,
)
from app.export.registry import ExportPayload, ExportSection, build_registry

DEFAULTS = dict(max_files=100, max_total_bytes=100 * 1024 * 1024, max_single_file_bytes=25 * 1024 * 1024)

SAMPLE = ExportPayload(
    title="Rate Sheet",
    source_name="rates",
    content="",
    headers=["item", "rate"],
    rows=[["checking", "0.1"], ["savings", "0.2"]],
    sections=[],
)


def test_available_formats_match_live_registry() -> None:
    built = build_registry()
    assert built.get("native") is not None
    assert built.get("docx") is not None
    assert built.get("xlsx") is not None
    assert built.get("pdf") is not None


def test_text_and_markdown_exporters() -> None:
    text = TextExporter().render(SAMPLE)
    assert b"Rate Sheet" in text
    markdown = MarkdownExporter().render(SAMPLE)
    assert isinstance(markdown, bytes)
    assert b"Rate Sheet" in markdown
    assert b"# Rate Sheet" in markdown


def test_markdown_artifacts_stripped_from_documents() -> None:
    payload = ExportPayload(
        title="repo",
        sections=[
            ExportSection(
                heading="Notes",
                body=(
                    "**Important:** reads `two` tables.\n\n"
                    "***\n\n"
                    "## subsection\n\n"
                    "* first\n* second\n\n"
                    "This *is* fine."
                ),
            )
        ],
    )
    rendered = TextExporter().render(payload).decode("utf-8")
    assert "**" not in rendered
    assert "***" not in rendered
    assert "`" not in rendered
    assert "#" not in rendered
    assert "Important: reads two tables." in rendered
    assert "is" in rendered
    assert "- first" in rendered


def test_text_table_renders_aligned_columns() -> None:
    from app.export.exporters.document_exporters import _format_text_table

    table = _format_text_table(
        ["File", "Signature"],
        [["src/main.py:2", "def main()"], ["src/cli.py:12", "async def run()"]],
    )
    lines = table.splitlines()
    assert "File" in lines[0]
    assert lines[1].startswith("  --")
    # data rows keep the same leading offset -> aligned columns
    assert lines[2].startswith("  src/main.py:2")
    assert lines[3].startswith("  src/cli.py:12")


def test_docx_renders_section_tables() -> None:
    pytest.importorskip("docx")
    payload = ExportPayload(
        title="repo",
        sections=[
            ExportSection(
                heading="API Endpoints",
                body="fallback",
                headers=["File", "Definition"],
                rows=[["src/app.py", '@app.get("/health")']],
            )
        ],
    )
    data = DocxExporter().render(payload)
    archive = zipfile.ZipFile(BytesIO(data))
    xml = archive.read("word/document.xml").decode("utf-8")
    assert "<w:tbl>" in xml
    assert 'w:t="File"' in xml or '>File</w:t>' in xml or "File" in xml
    assert "@app.get" in xml


def test_pdf_renders_section_tables() -> None:
    pytest.importorskip("reportlab")
    payload = ExportPayload(
        title="repo",
        sections=[
            ExportSection(
                heading="API Endpoints",
                body="fallback",
                headers=["File", "Definition"],
                rows=[["src/app.py", '@app.get("/health")']],
            )
        ],
    )
    data = PdfExporter().render(payload)
    assert data.startswith(b"%PDF")


def test_json_exporter_round_trips() -> None:
    data = json.loads(JsonExporter().render(SAMPLE).decode("utf-8"))
    assert data["title"] == "Rate Sheet"
    assert data["rows"][0] == ["checking", "0.1"]


def test_csv_exporter_contains_cells() -> None:
    csv_bytes = CsvExporter().render(SAMPLE)
    assert b"item,rate" in csv_bytes
    assert b"checking" in csv_bytes


def test_html_exporter_contains_markup() -> None:
    rendered = HtmlExporter().render(SAMPLE).decode("utf-8")
    assert "<h1>Rate Sheet</h1>" in rendered
    assert "</table>" in rendered


def test_docx_exporter_produces_openable_document() -> None:
    from docx import Document

    document = Document(BytesIO(DocxExporter().render(SAMPLE)))
    assert any("Rate Sheet" in paragraph.text for paragraph in document.paragraphs)


def test_xlsx_exporter_produces_openable_spreadsheet() -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(XlsxExporter().render(SAMPLE)))
    sheet = workbook.active
    assert sheet["A1"].value == "item"
    assert sheet["B1"].value == "rate"
    assert sheet.max_row >= 2


def test_pdf_exporter_produces_pdf_header() -> None:
    rendered = (
        build_registry().get("pdf").render(SAMPLE)  # type: ignore[union-attr]
    )
    assert rendered.startswith(b"%PDF")


def test_native_exporter_passthrough() -> None:
    payload = ExportPayload(metadata={"native_bytes": b"raw original bytes"})
    assert build_registry().get("native").render(payload) == b"raw original bytes"  # type: ignore[union-attr]


def test_native_exporter_requires_bytes() -> None:
    from app.export.registry import NativeExporter

    with pytest.raises(ValueError):
        NativeExporter().render(ExportPayload())


def test_safe_zip_entry_keeps_folder_structure() -> None:
    assert safe_relpath("uploaded_documents/policy.pdf") == "uploaded_documents/policy.pdf"
    assert safe_relpath("github/acme/api/source/app/main.py") == "github/acme/api/source/app/main.py"


def test_safe_zip_entry_defeats_traversal() -> None:
    entry = safe_relpath("uploaded_documents/../../etc/passwd")
    assert ".." not in entry
    assert entry.startswith("uploaded_documents/")


def test_write_zip_to_path_contains_manifest() -> None:
    tmpdir = pytest.importorskip("tempfile").gettempdir()
    import os

    target = os.path.join(tmpdir, "export-neo-test.zip")
    entries = [
        ZipEntry(name="uploaded_documents/policy.pdf", data=b"%PDF"),
    ]
    write_zip_to_path(target, entries, **DEFAULTS)
    try:
        with zipfile.ZipFile(target) as archive:
            assert archive.namelist()[0] == "manifest.json"
            manifest = json.loads(archive.read("manifest.json"))
            assert manifest["file_count"] == 1
            assert archive.read("uploaded_documents/policy.pdf") == b"%PDF"
    finally:
        os.remove(target)


def test_write_zip_enforces_total_size_limit() -> None:
    tmpdir = pytest.importorskip("tempfile").gettempdir()
    import os

    target = os.path.join(tmpdir, "export-neo-big.zip")
    entries = [ZipEntry(name=f"file-{i}.txt", data=b"x" * 1024) for i in range(120)]
    with pytest.raises(ExportLimitError):
        write_zip_to_path(target, entries, max_files=100, max_total_bytes=100 * 1024, max_single_file_bytes=1024 * 1024)
    assert not os.path.exists(target)


# --- Architecture diagram -----------------------------------------------------

_DIAGRAM_LAYERS = [
    ArchitectureLayer("Entry Point", ["MainApplication"]),
    ArchitectureLayer("API / REST Layer", ["ProductController", "OrderController"]),
    ArchitectureLayer("Business Logic / Services", ["ProductService"]),
    ArchitectureLayer("Data Access", ["ProductRepository"]),
    ArchitectureLayer("Data Models", []),
    ArchitectureLayer("Config & Infrastructure", ["AppConfig"]),
]


def test_architecture_text_diagram_lists_observed_components() -> None:
    text = render_architecture_text(_DIAGRAM_LAYERS)
    for marker in (
        "External clients",
        "ENTRY POINT",
        "API / REST LAYER",
        "BUSINESS LOGIC / SERVICES",
        "DATA ACCESS",
        "CONFIG & INFRASTRUCTURE",
        "MainApplication",
        "ProductController",
        "ProductRepository",
        "AppConfig",
        "Data / persistence",
    ):
        assert marker in text
    assert "(none observed)" in text  # the empty Data Models layer


def test_architecture_png_renders_png_bytes() -> None:
    pytest.importorskip("PIL")
    png = render_architecture_png(_DIAGRAM_LAYERS)
    assert png is not None
    assert png.startswith(b"\x89PNG")
    assert len(png) > 100


def test_architecture_png_skips_empty_layers() -> None:
    pytest.importorskip("PIL")
    assert render_architecture_png([ArchitectureLayer("Entry Point", [])]) is None


def test_docx_embeds_architecture_png() -> None:
    pytest.importorskip("docx")
    payload = ExportPayload(
        title="repo",
        diagram_png=render_architecture_png(_DIAGRAM_LAYERS),
        sections=[
            ExportSection("System Architecture", "fake caption", "Inferred from sampled source files"),
        ],
    )
    data = DocxExporter().render(payload)
    assert isinstance(data, bytes)
    archive = zipfile.ZipFile(BytesIO(data))
    assert any(name.startswith("word/media/") for name in archive.namelist())


def test_pdf_embeds_architecture_png() -> None:
    pytest.importorskip("reportlab")
    payload = ExportPayload(
        title="repo",
        diagram_png=render_architecture_png(_DIAGRAM_LAYERS),
        sections=[
            ExportSection("System Architecture", "fake caption", "Inferred from sampled source files"),
        ],
    )
    data = PdfExporter().render(payload)
    assert data.startswith(b"%PDF")


def test_manifest_summarizes_sources() -> None:
    manifest = build_manifest([])
    assert manifest["file_count"] == 0
    assert manifest["entries"] == []
